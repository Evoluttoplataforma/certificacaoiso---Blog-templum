# -*- coding: utf-8 -*-
import glob
import os
import re

import openpyxl

BASE = r"C:\Users\Raquel\Desktop\Marketing Templum\SemRush"

HEAD = [
    "certificação iso 9001",
    "certificacao iso 9001",
    "certificação iso 27001",
    "certificacao iso 27001",
    "certificação iso 14001",
    "certificacao iso 14001",
    "certificação iso 45001",
    "certificacao iso 45001",
    "quanto custa certificação iso 9001",
    "quanto custa certificacao iso 9001",
    "quanto custa certificação iso 27001",
    "como conseguir a certificação iso 14001",
    "como certificar iso 14001",
]


def norm(s):
    if not s:
        return ""
    s = str(s).lower().strip()
    for a, b in (
        ("ç", "c"),
        ("ã", "a"),
        ("á", "a"),
        ("à", "a"),
        ("é", "e"),
        ("ê", "e"),
        ("í", "i"),
        ("ó", "o"),
        ("ô", "o"),
        ("ú", "u"),
    ):
        s = s.replace(a, b)
    return re.sub(r"\s+", " ", s)


def load_positions(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    headers = None
    by_kw = {}
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            headers = row
            continue
        if not row or not row[0]:
            continue
        r = dict(zip(headers, row))
        k = norm(r.get("Keyword"))
        if k not in by_kw or (r.get("Position") or 999) < (by_kw[k].get("Position") or 999):
            by_kw[k] = r
    return by_kw


def main():
    pos_file = BASE + r"\certificacaoiso.com.br-organic.Positions-br-20261011-2026-10-12T14_13_52Z.xlsx"
    by_kw = load_positions(pos_file)

    print("=== Posicoes organicas (head terms) ===")
    for h in HEAD:
        r = by_kw.get(norm(h))
        if r:
            print(
                f"{h}: pos {r.get('Position')} vol {r.get('Search Volume')} kd {r.get('Keyword Difficulty')} -> {r.get('URL')}"
            )
        else:
            print(f"{h}: nao rankeia no export")

    print("\n=== certificacao iso* pos <= 30 (melhor URL por keyword) ===")
    hits = []
    for k, r in by_kw.items():
        if "certificacao iso" not in k:
            continue
        p = r.get("Position") or 999
        if p <= 30:
            hits.append(
                (
                    p,
                    r.get("Search Volume") or 0,
                    r.get("Keyword"),
                    r.get("URL"),
                )
            )
    hits.sort()
    for h in hits:
        print(h)

    # Keyword magic exports (multiple files = different seed?)
    kw_files = sorted(glob.glob(BASE + r"\certificação-iso_all-keywords_br_2026-10-12*.xlsx"))
    print(f"\n=== Keyword exports ({len(kw_files)} files) ===")
    for f in kw_files:
        wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
        ws = wb.active
        hdr = next(ws.iter_rows(max_row=1, values_only=True))
        print(os.path.basename(f), "| cols:", hdr[:6])
        # find rows matching head
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not row[0]:
                continue
            kw = norm(row[0])
            if kw in {norm(h) for h in HEAD}:
                print("  ", row[:8])

    serp_files = glob.glob(BASE + r"\certificação-iso_serp_urls_br_2026-10-12*.xlsx")
    if serp_files:
        print("\n=== SERP URLs (sample) ===", os.path.basename(serp_files[0]))
        wb = openpyxl.load_workbook(serp_files[0], read_only=True, data_only=True)
        ws = wb.active
        hdr = next(ws.iter_rows(max_row=1, values_only=True))
        print("cols:", hdr)
        n = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row:
                continue
            kw_raw = row[0] if row else ""
            if norm(kw_raw) in {norm(h) for h in HEAD[:8]}:
                print("KW:", kw_raw)
                for cell in row[1:11]:
                    if cell:
                        print(" ", cell)
                n += 1
                if n >= 4:
                    break


if __name__ == "__main__":
    main()
